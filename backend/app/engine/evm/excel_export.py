"""Executive EVM Excel export — ported close to verbatim from the reference
project's `backend/app/engine/evm/excel_export.py` (293 lines), using
`openpyxl` (approved as a new backend dependency for this slice — see
requirements.txt). Builds a styled, multi-sheet workbook with live SPI/CPI
formulas (recalculate if the user edits PV/EV/AC cells) and an embedded
S-curve line chart.

This is a downloadable Excel artifact, not a web page, so `DESIGN.md`'s
"never introduce new colors" rule (which governs the app's own UI) doesn't
apply here — the reference's own executive color palette is kept as-is,
which happens to already share our `--accent` amber (#F59E0B).
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0F1820"
AMBER = "F59E0B"
CERULEAN = "0F6B8A"
GREEN = "16A34A"
RED = "C0130A"
WHITE = "FFFFFF"
CREAM = "F7F5F1"
GRAY = "526070"
LIGHT = "E2DED8"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _align(h: str = "left", v: str = "center", wrap: bool = False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _border_bottom(color: str = LIGHT, style: str = "thin") -> Border:
    return Border(bottom=Side(style=style, color=color))


def _index_color(value, threshold: float = 1.0) -> str:
    if value is None:
        return GRAY
    if value >= threshold * 0.95:
        return GREEN
    if value >= threshold * 0.85:
        return "D97706"
    return RED


def build_evm_excel(project: dict, snapshots: list[dict], activities: list[dict]) -> bytes:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Executive Dashboard ────────────────────────────────────────
    ws = wb.active
    ws.title = "Executive Dashboard"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

    ws.column_dimensions["A"].width = 2
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 18
    for col in "KLMNOP":
        ws.column_dimensions[col].width = 15

    for row in range(1, 4):
        for col in range(1, 18):
            ws.cell(row=row, column=col).fill = _fill(NAVY)

    c = ws["B2"]
    c.value = f"EVM EXECUTIVE DASHBOARD — {project['name'].upper()}"
    c.font = Font(name="Barlow Condensed", size=18, bold=True, color=AMBER)
    c.alignment = _align("left")
    ws.merge_cells("B2:J2")
    ws.row_dimensions[2].height = 28

    c = ws["B3"]
    c.value = (
        f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}  |  "
        f"Baseline: {project.get('baseline_version', 'Target-1')}  |  "
        f"BAC: {project.get('bac', 0):,.0f} MH  |  "
        f"Period: {project.get('target_start', '')} → {project.get('target_end', '')}"
    )
    c.font = Font(name="Space Mono", size=8, color="8FAEC4", italic=True)
    ws.merge_cells("B3:P3")

    ws.row_dimensions[4].height = 10
    ws.row_dimensions[5].height = 14
    ws.row_dimensions[6].height = 32
    ws.row_dimensions[7].height = 24

    latest = snapshots[-1] if snapshots else {}
    kpis = [
        ("SPI", latest.get("spi"), "Schedule Perf. Index", 1.0),
        ("CPI", latest.get("cpi"), "Cost Perf. Index", 1.0),
        ("EV", latest.get("ev_cumulative"), "Earned Value (MH)", None),
        ("AC", latest.get("ac_cumulative"), "Actual Cost (MH)", None),
        ("PV", latest.get("pv_cumulative"), "Planned Value (MH)", None),
        ("EAC", latest.get("eac"), "Est. at Completion", None),
        ("TCPI", latest.get("tcpi"), "To-Complete Perf.", 1.0),
        ("CV", latest.get("cv"), "Cost Variance (MH)", None),
    ]

    col_start = 2
    for i, (code, val, label, threshold) in enumerate(kpis):
        col = col_start + i

        c_code = ws.cell(row=5, column=col)
        c_code.value = code
        c_code.font = Font(name="Barlow Condensed", size=13, bold=True, color=GRAY)
        c_code.alignment = _align("center")

        c_val = ws.cell(row=6, column=col)
        if val is not None:
            c_val.value = round(float(val), 2)
            c_val.number_format = "0.00" if threshold else "#,##0.0"
        else:
            c_val.value = "—"

        idx_color = _index_color(float(val), threshold) if val is not None and threshold else NAVY
        c_val.font = Font(
            name="Space Mono", size=14, bold=True,
            color=WHITE if val is not None and threshold and float(val) < threshold * 0.85 else NAVY,
        )
        c_val.fill = _fill(idx_color if threshold and val is not None else "F7F5F1")
        c_val.alignment = _align("center")
        c_val.border = Border(top=Side(style="medium", color=AMBER), bottom=Side(style="thin", color=LIGHT))

        c_lbl = ws.cell(row=7, column=col)
        c_lbl.value = label
        c_lbl.font = Font(name="Aptos", size=8, color=GRAY, italic=True)
        c_lbl.alignment = _align("center", wrap=True)

    ws.row_dimensions[8].height = 12

    c = ws["B9"]
    c.value = "S-CURVE DATA TABLE"
    c.font = Font(name="Barlow Condensed", size=12, bold=True, color=NAVY)
    c.fill = _fill("F0EDE8")
    c.alignment = _align("left")
    ws.merge_cells("B9:J9")
    ws.row_dimensions[9].height = 20

    headers = ["Date", "PV (MH)", "EV (MH)", "AC (MH)", "SPI*", "CPI*", "SV (MH)", "CV (MH)", "% Plan", "% Earned", "EAC", "TCPI*"]
    for j, h in enumerate(headers, 2):
        c = ws.cell(row=10, column=j)
        c.value = h
        c.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        c.fill = _fill(NAVY)
        c.alignment = _align("center")
        if "*" in h:
            c.font = Font(name="Space Mono", size=9, bold=True, color=AMBER)

    ws.row_dimensions[10].height = 20

    for i, snap in enumerate(snapshots, 11):
        row_fill = _fill(CREAM) if i % 2 == 0 else _fill(WHITE)
        row_data = [
            snap.get("snapshot_date"),
            snap.get("pv_cumulative"),
            snap.get("ev_cumulative"),
            snap.get("ac_cumulative"),
            None,  # SPI — live formula injected below
            None,  # CPI — live formula injected below
            snap.get("sv"),
            snap.get("cv"),
            snap.get("percent_complete_planned"),
            snap.get("percent_complete_earned"),
            snap.get("eac"),
            snap.get("tcpi"),
        ]
        for j, val in enumerate(row_data, 2):
            c = ws.cell(row=i, column=j)
            if val is not None:
                if isinstance(val, float):
                    c.value = round(val, 2)
                    c.number_format = "#,##0.00"
                else:
                    c.value = val
            c.fill = row_fill
            c.font = Font(name="Space Mono" if j in (2, 5, 6, 12) else "Aptos", size=9)
            c.alignment = _align("right" if j > 2 else "left")
            c.border = _border_bottom()

        pv_col = get_column_letter(3)
        ev_col = get_column_letter(4)
        ac_col = get_column_letter(5)
        ws.cell(row=i, column=6).value = f'=IF({pv_col}{i}>0,{ev_col}{i}/{pv_col}{i},"—")'
        ws.cell(row=i, column=6).font = Font(name="Space Mono", size=9, color=CERULEAN, bold=True)
        ws.cell(row=i, column=7).value = f'=IF({ac_col}{i}>0,{ev_col}{i}/{ac_col}{i},"—")'
        ws.cell(row=i, column=7).font = Font(name="Space Mono", size=9, color=GREEN, bold=True)

        tcpi_val = snap.get("tcpi")
        if tcpi_val is not None:
            tcpi_color = RED if tcpi_val > 1.10 else ("D97706" if tcpi_val > 1.0 else GREEN)
            ws.cell(row=i, column=13).font = Font(name="Space Mono", size=9, bold=True, color=tcpi_color)

    if snapshots:
        chart = LineChart()
        chart.title = "S-Curve: Planned vs Earned vs Actual"
        chart.style = 2
        chart.y_axis.title = "Cumulative Manhours"
        chart.x_axis.title = "Date"
        chart.width = 26
        chart.height = 14

        last_row = 10 + len(snapshots)
        data = Reference(ws, min_col=3, max_col=5, min_row=10, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.series[0].graphicalProperties.line.solidFill = CERULEAN
        chart.series[0].graphicalProperties.line.width = 22860
        chart.series[0].smooth = True
        chart.series[1].graphicalProperties.line.solidFill = GREEN
        chart.series[1].graphicalProperties.line.width = 22860
        chart.series[1].smooth = True
        chart.series[2].graphicalProperties.line.solidFill = AMBER
        chart.series[2].graphicalProperties.line.width = 22860
        chart.series[2].smooth = True
        ws.add_chart(chart, "N10")

    # ── Sheet 2: Baseline Activities ─────────────────────────────────────
    ws2 = wb.create_sheet("Baseline Activities")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 2

    act_headers = ["Activity Code", "Activity Name", "WBS", "Planned MH", "Baseline Start", "Baseline End"]
    col_widths = [18, 40, 20, 14, 16, 16]
    for j, (h, w) in enumerate(zip(act_headers, col_widths), 2):
        ws2.column_dimensions[get_column_letter(j)].width = w
        c = ws2.cell(row=1, column=j)
        c.value = h
        c.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        c.fill = _fill(NAVY)
        c.alignment = _align("center")
    ws2.row_dimensions[1].height = 20

    for i, act in enumerate(activities, 2):
        fill = _fill(CREAM) if i % 2 == 0 else _fill(WHITE)
        row = [
            act.get("task_code", ""),
            act.get("task_name", ""),
            act.get("wbs_code", ""),
            round(act.get("planned_manhours", 0), 2),
            act.get("baseline_start", ""),
            act.get("baseline_end", ""),
        ]
        for j, val in enumerate(row, 2):
            c = ws2.cell(row=i, column=j)
            c.value = val
            c.fill = fill
            c.font = Font(name="Space Mono" if j == 2 else "Aptos", size=9)
            c.alignment = _align("right" if j == 5 else "left")
            c.border = _border_bottom()

    last_act_row = 1 + len(activities) + 2
    note = ws2.cell(row=last_act_row, column=2)
    note.value = "* planned_manhours is immutable once this baseline is locked."
    note.font = Font(name="Aptos", size=8, italic=True, color=GRAY)
    ws2.merge_cells(f"B{last_act_row}:F{last_act_row}")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
